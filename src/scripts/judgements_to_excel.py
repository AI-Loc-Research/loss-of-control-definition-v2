"""Task 2C: one model's threat-model judgements (outputs/*.json) -> one Excel file.

    uv run python -m src.scripts.judgements_to_excel

Set the two paths below, then run from the repo root.

Sheet "judgements": one row per scenario, in the same order as output/scenarios_final.json
(= the phase-1 Excel). Columns: paper_id, scenario_id, scenario_text, then for each of the
7 components the 3 columns <component>, <component>_evidence, <component>_uncertainty
(21 columns, same names as in scenario_final_resolved.xlsx).

Sheet "summary": the per-scenario counts, as in summary.csv.

A scenario without a valid result (no file, or a parse-error placeholder) still gets a
row, with empty judgement cells, and is listed when the script runs.
"""

import glob
import json
import os

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# --- Files paths (relative to the repo root) ---
INPUT_PATH = "output/glm-5.3/07_task_2c/outputs" # folder with <paper_id>_<scenario_id>.json
SHEET_OUTPUT_PATH = "output/glm-5.3/07_task_2c/llm_glm5.3_scenario_judgement.xlsx"  # Excel file to write

SCENARIOS_PATH = "output/scenarios_final.json"   # scenario order and text
COMPONENTS = [
    "threat_source", "objective_or_harmful_outcome", "capability", "knowledge",
    "access", "constraints_or_enabling_conditions", "target_or_asset_at_risk",
]


def load_judgements(input_path: str) -> dict:
    """{"<paper_id>_<scenario_id>": parsed JSON} for every valid result file."""
    results = {}
    for path in glob.glob(os.path.join(input_path, "*.json")):
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        if data.get("parse_error"):
            continue   # placeholder for a failed answer, not a result
        results[os.path.basename(path)[:-5]] = data
    return results


def build_tables(judgements: dict, scenarios: list[dict]) -> tuple[pd.DataFrame, pd.DataFrame, list[str]]:
    rows, summary, missing = [], [], []
    for s in scenarios:
        key = f"{s['paper_id']}_{s['scenario_id']}"
        row = {"paper_id": s["paper_id"], "scenario_id": s["scenario_id"], "scenario_text": s["scenario_text"]}
        data = judgements.get(key)
        if data is None:
            missing.append(key)
        threat_model = (data or {}).get("threat_model", {})
        evidence = (data or {}).get("evidence", {})
        flags = (data or {}).get("uncertainty_flags", {})
        for comp in COMPONENTS:
            row[comp] = threat_model.get(comp, "") if data else ""
            row[f"{comp}_evidence"] = evidence.get(comp, "") if data else ""
            row[f"{comp}_uncertainty"] = flags.get(comp, "") if data else ""
        rows.append(row)
        if data:
            specified = sum(threat_model.get(c, "Not specified") != "Not specified" for c in COMPONENTS)
            summary.append({
                "paper_id": s["paper_id"], "scenario_id": s["scenario_id"],
                "num_specified_components": specified,
                "num_not_specified": len(COMPONENTS) - specified,
                "num_clear": sum(flags.get(c) == "Clear" for c in COMPONENTS),
                "num_ambiguous": sum(flags.get(c) == "Ambiguous" for c in COMPONENTS),
            })
    return pd.DataFrame(rows), pd.DataFrame(summary), missing


def format_sheet(ws, df: pd.DataFrame, wide_text: set[str]) -> None:
    """Bold frozen header, a filter on every column, readable column widths."""
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "C2"   # keep paper_id + scenario_id visible while scrolling
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, start=1):
        letter = get_column_letter(i)
        if col in wide_text:
            ws.column_dimensions[letter].width = 50
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        else:
            longest = max([len(str(col))] + [len(str(v)) for v in df[col]])
            ws.column_dimensions[letter].width = min(longest + 2, 30)


def main() -> None:
    if not os.path.isdir(INPUT_PATH):
        raise FileNotFoundError(f"Folder not found: {INPUT_PATH}")
    if not SHEET_OUTPUT_PATH.lower().endswith(".xlsx"):
        raise ValueError(f"SHEET_OUTPUT_PATH must end with .xlsx: {SHEET_OUTPUT_PATH}")
    with open(SCENARIOS_PATH, encoding="utf-8") as f:
        scenarios = json.load(f)

    judgements = load_judgements(INPUT_PATH)
    table, summary, missing = build_tables(judgements, scenarios)
    unknown = sorted(set(judgements) - {f"{s['paper_id']}_{s['scenario_id']}" for s in scenarios})

    out_dir = os.path.dirname(SHEET_OUTPUT_PATH)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    text_cols = {"scenario_text"} | {c for c in table.columns if c.endswith("_evidence")} | set(COMPONENTS)
    with pd.ExcelWriter(SHEET_OUTPUT_PATH, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="judgements", index=False)
        format_sheet(writer.sheets["judgements"], table, text_cols)
        summary.to_excel(writer, sheet_name="summary", index=False)
        format_sheet(writer.sheets["summary"], summary, set())

    print(f"{len(table) - len(missing)} of {len(table)} scenarios with judgements -> {SHEET_OUTPUT_PATH}")
    print(f"  sheet 'judgements': {len(table)} rows x {len(table.columns)} columns | sheet 'summary': {len(summary)} rows")
    if missing:
        print(f"  no valid result (empty row): {missing}")
    if unknown:
        print(f"  result files not in {SCENARIOS_PATH} (ignored): {unknown}")


if __name__ == "__main__":
    main()
