"""Task 2C: merge several models' threat-model judgements into one Excel sheet.

    uv run python -m src.scripts.merge_judgements

Set MODELS and SHEET_OUTPUT_PATH below, then run from the repo root.

One row per scenario, in the order of output/scenarios_final.json (= the phase-1 Excel).
Columns: paper_id, scenario_id, scenario_text, then for each of the 7 components, for each
model, the 3 fields:
    threat_source_kimi_2.5, threat_source_evidence_kimi_2.5, threat_source_uncertainty_kimi_2.5,
    threat_source_glm_5.3, threat_source_evidence_glm_5.3, ... then objective_or_harmful_outcome_...

Each model is read from its source data:
    "excel": the phase-1 library Excel (Sheet1 columns <component>, <component>_evidence,
             <component>_uncertainty), e.g. Kimi K2.5
    "json":  a Task 2C outputs folder (<paper_id>_<scenario_id>.json), e.g. GLM, DeepSeek
A scenario a model has no valid result for gets empty cells and is listed when the script runs.
"""

import glob
import json
import os

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# --- Set these (paths relative to the repo root) ---
# (label used in the column names, source type, path)
MODELS = [
    ("kimi_2.5", "excel", "materials/scenario_final_resolved.xlsx"),
    ("glm_5.3", "json", "output/glm-5.3/07_task_2c/outputs"),
    ("deepseek_v4.1", "json", "output/deepseek-v4.1-flash/07_task_2c/outputs"),
]
SHEET_OUTPUT_PATH = "output/07_task_2c_llm_judgements_merged.xlsx"

SCENARIOS_PATH = "output/scenarios_final.json"   # scenario order and text
EXCEL_SHEET = "Sheet1"                           # sheet of the phase-1 Excel with the 193 scenarios
PAPER_ID_FIXES = {"barrett_2026": "barrett_2016"}  # known typo in the phase-1 Excel
COMPONENTS = [
    "threat_source", "objective_or_harmful_outcome", "capability", "knowledge",
    "access", "constraints_or_enabling_conditions", "target_or_asset_at_risk",
]
FIELDS = ("", "_evidence", "_uncertainty")   # value, evidence, uncertainty flag


def _cell(value) -> str:
    """Excel cell -> text ('' for an empty cell)."""
    return "" if pd.isna(value) else str(value)


def load_from_excel(path: str) -> dict:
    """{"<paper_id>_<scenario_id>": {"<component><field>": text}} from the phase-1 Excel."""
    df = pd.read_excel(path, sheet_name=EXCEL_SHEET, dtype=object)
    wanted = [comp + field for comp in COMPONENTS for field in FIELDS]
    missing = [c for c in wanted if c not in df.columns]
    if missing:
        raise ValueError(f"{path}: columns not found: {missing}")
    results = {}
    for _, row in df.iterrows():
        paper, scen = _cell(row["paper_id"]).strip(), _cell(row["scenario_id"]).strip()
        if not paper or not scen:
            continue   # empty rows after the data
        paper = PAPER_ID_FIXES.get(paper, paper)
        results[f"{paper}_{scen}"] = {c: _cell(row[c]) for c in wanted}
    return results


def load_from_json(folder: str) -> dict:
    """{"<paper_id>_<scenario_id>": {"<component><field>": text}} from a 2C outputs folder."""
    if not os.path.isdir(folder):
        raise FileNotFoundError(f"Folder not found: {folder}")
    results = {}
    for path in glob.glob(os.path.join(folder, "*.json")):
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        if data.get("parse_error"):
            continue   # placeholder for a failed answer, not a result
        parts = {"": data.get("threat_model", {}), "_evidence": data.get("evidence", {}),
                 "_uncertainty": data.get("uncertainty_flags", {})}
        results[os.path.basename(path)[:-5]] = {
            comp + field: _cell(parts[field].get(comp)) for comp in COMPONENTS for field in FIELDS}
    return results


def build_table(scenarios: list[dict], judgements: dict) -> tuple[pd.DataFrame, dict]:
    """One row per scenario; columns ordered by component, then model, then field."""
    rows, missing = [], {label: [] for label in judgements}
    for s in scenarios:
        key = f"{s['paper_id']}_{s['scenario_id']}"
        row = {"paper_id": s["paper_id"], "scenario_id": s["scenario_id"], "scenario_text": s["scenario_text"]}
        for label, results in judgements.items():
            if key not in results:
                missing[label].append(key)
        for comp in COMPONENTS:
            for label, results in judgements.items():
                for field in FIELDS:
                    row[f"{comp}{field}_{label}"] = results.get(key, {}).get(comp + field, "")
        rows.append(row)
    return pd.DataFrame(rows), missing


def format_sheet(ws, df: pd.DataFrame) -> None:
    """Bold frozen header, a filter on every column, wrapped long text."""
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "D2"   # keep paper_id, scenario_id, scenario_text visible
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, start=1):
        letter = get_column_letter(i)
        if col in ("paper_id", "scenario_id") or "_uncertainty_" in col:
            ws.column_dimensions[letter].width = max(14, min(len(col) + 2, 30))
        else:
            ws.column_dimensions[letter].width = 45
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    if not SHEET_OUTPUT_PATH.lower().endswith(".xlsx"):
        raise ValueError(f"SHEET_OUTPUT_PATH must end with .xlsx: {SHEET_OUTPUT_PATH}")
    with open(SCENARIOS_PATH, encoding="utf-8") as f:
        scenarios = json.load(f)
    loaders = {"excel": load_from_excel, "json": load_from_json}
    judgements = {label: loaders[kind](path) for label, kind, path in MODELS}

    table, missing = build_table(scenarios, judgements)
    out_dir = os.path.dirname(SHEET_OUTPUT_PATH)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    with pd.ExcelWriter(SHEET_OUTPUT_PATH, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="judgements", index=False)
        format_sheet(writer.sheets["judgements"], table)

    print(f"{len(table)} scenarios x {len(table.columns)} columns -> {SHEET_OUTPUT_PATH}")
    known = {f"{s['paper_id']}_{s['scenario_id']}" for s in scenarios}
    for label, results in judgements.items():
        line = f"  {label:14} {len(table) - len(missing[label])} of {len(table)} scenarios"
        if missing[label]:
            line += f" | no result (empty cells): {missing[label]}"
        extra = sorted(set(results) - known)
        if extra:
            line += f" | not in {SCENARIOS_PATH} (ignored): {extra}"
        print(line)


if __name__ == "__main__":
    main()
